"""Local storage and document conversion fixtures; no model or private data access."""
import hashlib
import io
import json
from pathlib import Path

from burmese_reader import annotations, debug_pos, pdf_cache, settings, srs, text_extraction
from burmese_reader.application import create_app


def test_annotation_storage_round_trip_and_clear(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "ANNOTATIONS_PATH", tmp_path / "annotations.json")
    app = create_app(load_resources=False)
    with app.app_context():
        annotations.set_annotation("စာ", "a note\n")
        assert annotations.get_annotation("စာ") == "a note"
        annotations.state.ANNOTATIONS.clear()
        annotations.load_annotations()
        assert annotations.get_annotation("စာ") == "a note"
        annotations.set_annotation("စာ", "")
        assert annotations.get_annotation("စာ") == ""
    assert json.loads(settings.ANNOTATIONS_PATH.read_text()) == {}


def test_srs_serialization_and_review_schedule(tmp_path, monkeypatch):
    # Explicit opt-in for this fixture; public routes remain disabled.
    monkeypatch.setattr(srs, "READING_SRS_SAVE_ENABLED", True)
    store = srs.ReadingSRS(tmp_path / "srs.json")
    store.observe_tokens(["စာ", "စာ", "ဖတ်"])
    assert store.get_next_card()["head"] == "စာ"
    store.grade_card("စာ", True)
    first = store.cards["စာ"]
    assert first.stats.total_seen == 2
    assert first.review.repetitions == 1
    assert first.review.due > first.stats.first_seen
    again = srs.ReadingSRS(store.state_path)
    again._ensure_loaded()
    assert again.cards["စာ"].to_dict() == first.to_dict()


def test_unicode_text_import_keeps_line_breaks_and_pages():
    source = "စာ\n" * 1100
    pages = text_extraction._extract_text_pages_simple(source)
    assert len(pages) > 1
    assert "\n".join(pages) == source
    app = create_app(load_resources=False)
    response = app.test_client().post(
        "/api/extract_text_simple", data={"file": (io.BytesIO(source.encode()), "fixture.txt")}
    )
    assert response.status_code == 200
    assert response.json["ok"]


def test_pdf_cache_closes_only_its_own_document():
    first = create_app(load_resources=False)
    second = create_app(load_resources=False)
    with first.app_context():
        token = pdf_cache._cache_pdf_to_disk(b"%PDF-fixture")
        path = Path(pdf_cache._get_cached_pdf_path(token))
    try:
        assert second.test_client().get(f"/api/serve_pdf/{token}").status_code == 410
        assert first.test_client().get(f"/api/serve_pdf/{token}").data == b"%PDF-fixture"
        assert first.test_client().post("/api/close_pdf", json={"cache_id": token}).status_code == 200
        assert not path.exists()
    finally:
        with first.app_context():
            pdf_cache._evict_cached_pdf(token)


def test_extracted_debug_template_preserves_rendered_bytes():
    fixture = json.loads((Path(__file__).parent / "fixtures/debug-template.json").read_text(encoding="utf-8"))
    rendered = debug_pos._render_pos_sentence_debug_html(fixture["input"])
    assert hashlib.sha256(rendered.encode()).hexdigest() == fixture["sha256"]


def test_legacy_named_imports_resolve_without_model_loading():
    from app import normalize_burmese_for_segmentation, segment_with_pipeline, app
    assert normalize_burmese_for_segmentation("စာ\u200b") == "စာ\u200b"
    assert callable(segment_with_pipeline)
    assert app.test_client().get("/ping").json["ok"]


def test_docx_text_and_pdf_geometry_from_synthetic_documents():
    import docx
    import fitz
    from burmese_reader.pdf_extraction import _extract_pdf_pages_and_dims_from_bytes

    document = docx.Document()
    document.add_paragraph("မြန်မာစာ")
    document.add_paragraph("ဖတ်")
    stream = io.BytesIO()
    document.save(stream)
    app = create_app(load_resources=False)
    client = app.test_client()
    response = client.post("/api/extract_text_simple", data={
        "file": (io.BytesIO(stream.getvalue()), "fixture.docx")
    })
    assert response.status_code == 200
    assert response.json["text"] == "မြန်မာစာ\nဖတ်"

    with fitz.open() as pdf:
        page = pdf.new_page(width=300, height=400)
        page.insert_text((40, 60), "Fixture document")
        data = pdf.tobytes()
    pages, dimensions, average = _extract_pdf_pages_and_dims_from_bytes(data)
    assert pages == ["Fixture document\n"]
    assert dimensions[0]["width"] == 300
    assert average["height"] == 400
    with app.app_context():
        token = pdf_cache._cache_pdf_to_disk(data)
    try:
        result = client.post("/api/pdf_page_text", json={"cache_id": token, "page": 0})
        assert result.status_code == 200
        assert result.json["raw_text"] == pages[0]
        assert len(result.json["words"]) == 2
        assert result.json["structured_blocks"]
    finally:
        client.post("/api/close_pdf", json={"cache_id": token})
